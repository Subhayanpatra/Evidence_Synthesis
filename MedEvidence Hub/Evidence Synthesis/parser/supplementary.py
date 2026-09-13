from __future__ import annotations

import io
import json
import os
import re
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import quote, unquote, urljoin, urlparse

import fitz
import pandas as pd
import requests
from bs4 import BeautifulSoup
from docx import Document
from lxml import etree
from pptx import Presentation
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from config import EMAIL


# This module follows the tested notebook workflow while retaining the three
# fields consumed by the AI_extract paper workflow and frontend.
DEFAULT_DOWNLOAD_ROOT = Path(
    os.getenv("SUPPLEMENTARY_DOWNLOAD_DIR", Path("data") / "supplementary_materials")
)
MAX_FILE_SIZE = int(
    os.getenv("SUPPLEMENTARY_MAX_DOWNLOAD_BYTES", str(200 * 1024 * 1024))
)
MAX_EXTRACTED_CHARACTERS = int(
    os.getenv("SUPPLEMENTARY_MAX_EXTRACTED_CHARACTERS", "2000000")
)
MAX_NESTED_ZIP_DEPTH = int(os.getenv("SUPPLEMENTARY_MAX_ZIP_DEPTH", "5"))
MAX_FILES_INSIDE_ZIP = int(os.getenv("SUPPLEMENTARY_MAX_ZIP_FILES", "1000"))
CHECKPOINT_VERSION = 2

HEADERS = {
    "User-Agent": f"AI-Extract-Supplementary/2.0 ({EMAIL})",
    "Accept": "*/*",
    "Accept-Language": "en-US,en;q=0.9",
}

SUPPORTED_EXTENSIONS = {
    ".avi",
    ".csv",
    ".doc",
    ".docx",
    ".htm",
    ".html",
    ".json",
    ".md",
    ".mov",
    ".mp4",
    ".pdf",
    ".ppt",
    ".pptx",
    ".tsv",
    ".txt",
    ".xls",
    ".xlsx",
    ".xml",
    ".zip",
}
VIDEO_EXTENSIONS = {".avi", ".mov", ".mp4"}
IMAGE_EXTENSIONS = {
    ".bmp",
    ".gif",
    ".jpeg",
    ".jpg",
    ".png",
    ".svg",
    ".tif",
    ".tiff",
    ".webp",
}
SUPPLEMENTARY_PATTERNS = (
    r"supp",
    r"suppl",
    r"supporting",
    r"mmc\d+",
    r"moesm\d+",
    r"[-_]s\d{3,}",
    r"[-_]sm\d*",
    r"[-_]esm\d*",
)

WORD_NAMESPACE = (
    "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
)
WORD_NS = {"w": WORD_NAMESPACE}


def _create_session() -> requests.Session:
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        status=3,
        backoff_factor=1,
        status_forcelist=[429, 500, 502, 503, 504],
        allowed_methods=["GET"],
    )
    session = requests.Session()
    session.mount("https://", HTTPAdapter(max_retries=retry))
    session.headers.update(HEADERS)
    return session


def normalize_pmcid(pmcid: Any) -> str:
    value = str(pmcid or "").strip().upper()
    if not value.startswith("PMC"):
        value = f"PMC{value}"
    if not value[3:].isdigit():
        raise ValueError(f"Invalid PMCID: {value}")
    return value


def _normalize_pmcid(pmcid: Any) -> str:
    try:
        return normalize_pmcid(pmcid)
    except ValueError:
        return ""


def clean_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def filename_from_url(url: str) -> str:
    path = unquote(urlparse(url).path)
    return PurePosixPath(path).name


def _safe_file_name(value: str) -> str:
    name = filename_from_url(str(value)) or Path(unquote(str(value))).name
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" .")


def _normalized_file_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", unquote(str(value)).casefold())


def _truncate_content(text: str) -> str:
    if len(text) <= MAX_EXTRACTED_CHARACTERS:
        return text
    return (
        text[:MAX_EXTRACTED_CHARACTERS]
        + f"\n\n[Content truncated at {MAX_EXTRACTED_CHARACTERS:,} characters.]"
    )


def looks_like_html(data: bytes) -> bool:
    beginning = data[:1000].lstrip().lower()
    return (
        beginning.startswith(b"<!doctype html")
        or beginning.startswith(b"<html")
        or b"<title>checking your browser" in beginning
        or b"recaptcha" in beginning
    )


def validate_download(filename: str, data: bytes, content_type: str = "") -> bool:
    if not data:
        raise ValueError("Downloaded file is empty")
    if len(data) > MAX_FILE_SIZE:
        raise ValueError("Downloaded file exceeds maximum size")

    extension = PurePosixPath(filename).suffix.lower()
    if looks_like_html(data) and extension not in {".html", ".htm"}:
        raise ValueError("Server returned an HTML page instead of the real file")
    if extension == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    if extension in {".docx", ".xlsx", ".pptx", ".zip"}:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise ValueError(f"Invalid {extension} file: ZIP signature not found")
    return True


def download_pmc_xml(session: requests.Session, pmcid: str) -> tuple[bytes, str]:
    xml_url = (
        "https://www.ebi.ac.uk/europepmc/webservices/rest/"
        f"{pmcid}/fullTextXML"
    )
    response = session.get(xml_url, timeout=(15, 90))
    response.raise_for_status()
    if looks_like_html(response.content):
        raise ValueError("XML endpoint returned HTML")
    return response.content, xml_url


def detect_supplementary_in_xml(xml_data: bytes | str | None) -> bool:
    if not xml_data:
        return False
    soup = BeautifulSoup(xml_data, "xml")
    if soup.find_all("supplementary-material"):
        return True

    for element in soup.find_all(["media", "ext-link"]):
        href = element.get("xlink:href") or element.get("href") or ""
        marker_text = " ".join(
            [
                str(element.get("id", "")),
                str(element.get("content-type", "")),
                href,
                element.get_text(" ", strip=True),
            ]
        ).lower()
        if any(
            marker in marker_text
            for marker in (
                "supplement",
                "supplementary",
                "additional file",
                "supporting information",
            )
        ):
            return True
    return False


def find_supplementary_files(
    xml_content: bytes | str | None,
) -> list[dict[str, str]]:
    """Compatibility helper: return supplementary file references from PMC XML."""
    if not xml_content:
        return []
    soup = BeautifulSoup(xml_content, "xml")
    candidates: list[str] = []

    for node in soup.find_all(["supplementary-material", "inline-supplementary-material"]):
        for child in node.find_all(href=True):
            candidates.append(str(child.get("href", "")))
        for child in node.find_all(attrs={"xlink:href": True}):
            candidates.append(str(child.get("xlink:href", "")))

    for element in soup.find_all(["media", "ext-link", "a"]):
        href = str(element.get("xlink:href") or element.get("href") or "")
        file_name = filename_from_url(href)
        if file_name and any(
            re.search(pattern, file_name, re.IGNORECASE)
            for pattern in SUPPLEMENTARY_PATTERNS
        ):
            candidates.append(href)

    unique: dict[str, dict[str, str]] = {}
    for href in candidates:
        file_name = _safe_file_name(href)
        if not file_name or Path(file_name).suffix.casefold() not in SUPPORTED_EXTENSIONS:
            continue
        unique[_normalized_file_name(file_name)] = {
            "File_Name": file_name,
            "Href": unquote(str(href)).strip(),
        }
    return list(unique.values())


def extract_supplementary_links_from_html(
    html: str, article_url: str
) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "lxml")
    collected_links: dict[str, dict[str, str]] = {}

    def add_links(container: Any) -> None:
        for anchor in container.find_all("a", href=True):
            absolute_url = urljoin(article_url, anchor.get("href", "").strip())
            if "/articles/instance/" in absolute_url and "/bin/" in absolute_url:
                filename = filename_from_url(absolute_url)
                if filename:
                    collected_links[absolute_url] = {
                        "url": absolute_url,
                        "filename": filename,
                        "link_text": clean_text(anchor.get_text(" ", strip=True)),
                    }

    for selector in (
        "section.sm",
        "section.supplementary-material",
        "div.supplementary-material",
        "[id^='SM']",
        "[id^='SD']",
    ):
        for container in soup.select(selector):
            add_links(container)

    heading_pattern = re.compile(
        r"supplementary|supporting information|additional files?", re.IGNORECASE
    )
    for heading in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6"]):
        if not heading_pattern.search(heading.get_text(" ", strip=True)):
            continue
        parent_section = heading.find_parent("section")
        if parent_section:
            add_links(parent_section)
        for sibling in heading.find_next_siblings():
            if sibling.name in ["h1", "h2", "h3", "h4", "h5", "h6"]:
                break
            add_links(sibling)

    if not collected_links:
        for anchor in soup.find_all("a", href=True):
            absolute_url = urljoin(article_url, anchor.get("href", ""))
            if "/articles/instance/" in absolute_url and "/bin/" in absolute_url:
                filename = filename_from_url(absolute_url)
                if filename:
                    collected_links[absolute_url] = {
                        "url": absolute_url,
                        "filename": filename,
                        "link_text": clean_text(anchor.get_text(" ", strip=True)),
                    }
    return list(collected_links.values())


def download_pmc_html(session: requests.Session, pmcid: str) -> tuple[str, str]:
    article_url = f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/"
    response = session.get(
        article_url,
        headers={**HEADERS, "Referer": "https://pmc.ncbi.nlm.nih.gov/"},
        timeout=(15, 90),
    )
    response.raise_for_status()
    return response.text, article_url


def word_xml_text(element: etree._Element) -> str:
    parts: list[str] = []
    for node in element.iter():
        node_type = etree.QName(node).localname
        if node_type == "t" and node.text:
            parts.append(node.text)
        elif node_type == "tab":
            parts.append("\t")
        elif node_type in {"br", "cr"}:
            parts.append("\n")
    return clean_text("".join(parts))


def iter_word_blocks(parent: etree._Element):
    for child in parent:
        child_type = etree.QName(child).localname
        if child_type in {"p", "tbl"}:
            yield child
        else:
            yield from iter_word_blocks(child)


def extract_docx_raw_xml(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("File is not a genuine DOCX container")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if "word/document.xml" not in archive.namelist():
            raise ValueError("word/document.xml is missing")
        root = etree.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", namespaces=WORD_NS)
    if body is None:
        raise ValueError("Word document body is missing")

    output: list[str] = []
    table_number = 0
    for block in iter_word_blocks(body):
        block_type = etree.QName(block).localname
        if block_type == "p":
            paragraph_text = word_xml_text(block)
            if paragraph_text:
                output.append(paragraph_text)
        elif block_type == "tbl":
            table_number += 1
            table_rows: list[str] = []
            for row in block.xpath("./w:tr", namespaces=WORD_NS):
                cells: list[str] = []
                for cell in row.xpath("./w:tc", namespaces=WORD_NS):
                    cell_paragraphs = [
                        word_xml_text(paragraph)
                        for paragraph in cell.xpath(".//w:p", namespaces=WORD_NS)
                    ]
                    cells.append(" ".join(text for text in cell_paragraphs if text))
                if any(cells):
                    table_rows.append(" | ".join(cells))
            if table_rows:
                output.append(f"[TABLE {table_number}]\n" + "\n".join(table_rows))
    return "\n\n".join(output).strip()


def extract_docx_text(data: bytes) -> str:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("File is not a genuine DOCX container")
    try:
        document = Document(io.BytesIO(data))
        output: list[str] = []
        paragraphs = [
            clean_text(paragraph.text)
            for paragraph in document.paragraphs
            if clean_text(paragraph.text)
        ]
        if paragraphs:
            output.append("\n".join(paragraphs))
        for table_number, table in enumerate(document.tables, start=1):
            rows = []
            for row in table.rows:
                cells = [clean_text(cell.text) for cell in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                output.append(f"[TABLE {table_number}]\n" + "\n".join(rows))
        extracted_text = "\n\n".join(output).strip()
        if extracted_text:
            return extracted_text
    except Exception:
        pass
    return extract_docx_raw_xml(data)


def extract_pdf_text(data: bytes) -> str:
    output: list[str] = []
    with fitz.open(stream=data, filetype="pdf") as document:
        for page_number, page in enumerate(document, start=1):
            page_text = page.get_text("text").strip()
            if page_text:
                output.append(f"[PAGE {page_number}]\n{page_text}")
    return "\n\n".join(output).strip()


def extract_excel_text(data: bytes) -> str:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    return "\n\n".join(
        f"[SHEET: {sheet_name}]\n{dataframe.fillna('').to_csv(index=False)}"
        for sheet_name, dataframe in sheets.items()
    ).strip()


def extract_pptx_text(data: bytes) -> str:
    presentation = Presentation(io.BytesIO(data))
    output: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = clean_text(shape.text)
                if text:
                    slide_parts.append(text)
            if getattr(shape, "has_table", False):
                table_rows = []
                for row in shape.table.rows:
                    cells = [clean_text(cell.text) for cell in row.cells]
                    if any(cells):
                        table_rows.append(" | ".join(cells))
                if table_rows:
                    slide_parts.append("\n".join(table_rows))
        if slide_parts:
            output.append(f"[SLIDE {slide_number}]\n" + "\n".join(slide_parts))
    return "\n\n".join(output).strip()


def decode_text_file(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return ""


def is_normal_article_figure(filename: str) -> bool:
    """Return whether a filename is a normal PMC article figure."""
    name = PurePosixPath(str(filename or "").replace("\\", "/")).name.casefold()
    path = PurePosixPath(name)
    if path.suffix not in IMAGE_EXTENSIONS:
        return False
    stem = path.stem
    return any(
        re.fullmatch(pattern, stem, flags=re.IGNORECASE)
        for pattern in (
            r"(?:.*[_\-.])?fig(?:ure)?[_\-.]?\d+[_\-.]?html",
            r"fig(?:ure)?[_\-.]?\d+",
            r"(?:gr|ga)[_\-.]?\d+",
        )
    )


def _is_unsafe_zip_member(member_name: str) -> bool:
    path = PurePosixPath(str(member_name or "").replace("\\", "/"))
    return path.is_absolute() or ".." in path.parts


def _is_metadata_zip_file(member_name: str) -> bool:
    normalized = str(member_name or "").replace("\\", "/").casefold()
    return (
        normalized.startswith("__macosx/")
        or PurePosixPath(normalized).name
        in {".ds_store", "desktop.ini", "thumbs.db"}
    )


def extract_text_from_file(filename: str, data: bytes, _zip_depth: int = 0) -> str:
    extension = PurePosixPath(filename).suffix.lower()
    if extension == ".docx":
        return extract_docx_text(data)
    if extension == ".pdf":
        return extract_pdf_text(data)
    if extension in {".xlsx", ".xls"}:
        return extract_excel_text(data)
    if extension == ".pptx":
        return extract_pptx_text(data)
    if extension in {".txt", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".md"}:
        return decode_text_file(data)
    if extension in VIDEO_EXTENSIONS:
        return f"[Video supplementary material downloaded successfully.]\nVideo file: {filename}"
    if extension in {".doc", ".ppt"}:
        return (
            f"[Legacy {extension.upper()[1:]} downloaded successfully, "
            "but automatic text extraction is not supported.]"
        )
    if extension == ".zip":
        return _extract_zip_bytes(data, current_depth=_zip_depth)
    return ""


def _extract_zip_bytes(data: bytes, current_depth: int = 0) -> str:
    if current_depth >= MAX_NESTED_ZIP_DEPTH:
        raise ValueError("Maximum nested ZIP depth exceeded")
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise ValueError("File is not a valid ZIP archive")

    output: list[str] = []
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > MAX_FILES_INSIDE_ZIP:
            raise ValueError(f"ZIP contains too many files: {len(members)}")

        for member in members:
            if (
                member.is_dir()
                or member.file_size <= 0
                or member.file_size > MAX_FILE_SIZE
                or _is_unsafe_zip_member(member.filename)
                or _is_metadata_zip_file(member.filename)
            ):
                continue
            filename = PurePosixPath(member.filename).name
            extension = PurePosixPath(filename).suffix.lower()
            if (
                not filename
                or is_normal_article_figure(filename)
                or extension not in SUPPORTED_EXTENSIONS
            ):
                continue
            try:
                nested_text = extract_text_from_file(
                    filename,
                    archive.read(member),
                    _zip_depth=current_depth + 1,
                )
            except Exception:
                # One corrupt or over-depth member must not discard text
                # successfully extracted from the rest of the archive.
                continue
            if nested_text:
                output.append(f"ARCHIVE FILE: {filename}\n{nested_text}")
    return "\n\n".join(output).strip()


def extract_supplementary_content(file_path: str | Path) -> str:
    """Compatibility helper used by tests and any direct project callers."""
    path = Path(file_path)
    try:
        data = path.read_bytes()
        if path.suffix.casefold() == ".zip":
            text = _extract_zip_bytes(data)
        else:
            text = extract_text_from_file(path.name, data)
        return _truncate_content(text)
    except Exception as error:
        return f"[Content extraction failed: {error}]"


def download_binary(
    session: requests.Session, url: str, referer: str | None = None
) -> tuple[bytes, str, str]:
    request_headers = dict(HEADERS)
    if referer:
        request_headers["Referer"] = referer
    response = session.get(
        url,
        headers=request_headers,
        timeout=(15, 120),
        allow_redirects=True,
    )
    response.raise_for_status()
    return response.content, response.url, response.headers.get("Content-Type", "")


def download_europe_pmc_zip(
    session: requests.Session, pmcid: str
) -> tuple[dict[str, dict[str, Any]], str | None]:
    url = (
        "https://www.ebi.ac.uk/europepmc/webservices/rest/"
        f"{pmcid}/supplementaryFiles"
    )
    response = session.get(url, timeout=(15, 120), allow_redirects=True)
    response.raise_for_status()
    if not zipfile.is_zipfile(io.BytesIO(response.content)):
        return {}, "Europe PMC did not return a ZIP file"

    extracted_files: dict[str, dict[str, Any]] = {}
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        for member in archive.infolist():
            if (
                member.is_dir()
                or member.file_size > MAX_FILE_SIZE
                or _is_unsafe_zip_member(member.filename)
                or _is_metadata_zip_file(member.filename)
            ):
                continue
            filename = PurePosixPath(member.filename).name
            if filename and not is_normal_article_figure(filename):
                extracted_files[filename] = {
                    "data": archive.read(member),
                    "source_url": url,
                    "content_type": "",
                    "archive_member": member.filename,
                }
    return extracted_files, None


def collect_biostudies_files(node: Any) -> list[str]:
    collected: list[str] = []
    if isinstance(node, dict):
        files = node.get("files", [])
        if isinstance(files, list):
            for file_record in files:
                if (
                    isinstance(file_record, dict)
                    and file_record.get("path")
                    and file_record.get("type", "file") == "file"
                ):
                    collected.append(str(file_record["path"]))
        for key, value in node.items():
            if key != "files":
                collected.extend(collect_biostudies_files(value))
    elif isinstance(node, list):
        for item in node:
            collected.extend(collect_biostudies_files(item))
    return list(dict.fromkeys(collected))


def download_biostudies_files(
    session: requests.Session, pmcid: str
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    accession = f"S-EPMC{pmcid.replace('PMC', '')}"
    api_url = f"https://www.ebi.ac.uk/biostudies/api/v1/studies/{accession}"
    response = session.get(api_url, timeout=(15, 90))
    response.raise_for_status()
    file_paths = collect_biostudies_files(response.json())
    downloaded_files: dict[str, dict[str, Any]] = {}
    errors: list[str] = []

    for file_path in file_paths:
        download_url = (
            "https://www.ebi.ac.uk/biostudies/files/"
            f"{accession}/{quote(file_path, safe='/')}"
        )
        filename = PurePosixPath(file_path).name
        try:
            data, final_url, content_type = download_binary(session, download_url)
            validate_download(filename, data, content_type)
            downloaded_files[filename] = {
                "data": data,
                "source_url": final_url,
                "content_type": content_type,
            }
        except Exception as error:
            errors.append(f"{filename}: {type(error).__name__}: {error}")
    return downloaded_files, errors


def extract_doi_from_pmc_html(html: str) -> str | None:
    soup = BeautifulSoup(html, "lxml")
    doi_meta = soup.find("meta", attrs={"name": "citation_doi"})
    if doi_meta and doi_meta.get("content"):
        return str(doi_meta["content"]).strip()
    return None


def extract_doi_from_pmc_xml(xml_content: bytes | str | None) -> str | None:
    if not xml_content:
        return None
    soup = BeautifulSoup(xml_content, "xml")
    doi_node = soup.find("article-id", attrs={"pub-id-type": "doi"})
    if doi_node:
        doi = clean_text(doi_node.get_text(" ", strip=True))
        if doi:
            return doi
    return None


def download_springer_supplementary_files(
    session: requests.Session,
    doi: str | None,
    supplementary_links: list[str],
) -> tuple[dict[str, dict[str, Any]], list[str]]:
    if not doi:
        return {}, ["Publisher fallback: DOI was not found"]
    if not doi.lower().startswith(("10.1038/", "10.1007/", "10.1186/", "10.1134/")):
        return {}, [f"Publisher fallback is not configured for DOI {doi}"]

    downloaded_files: dict[str, dict[str, Any]] = {}
    errors: list[str] = []
    encoded_doi = quote(doi, safe="")
    for pmc_link in supplementary_links:
        filename = filename_from_url(pmc_link)
        if not filename:
            continue
        publisher_url = (
            "https://static-content.springer.com/esm/"
            f"art%3A{encoded_doi}/MediaObjects/{quote(filename, safe='._-')}"
        )
        try:
            data, final_url, content_type = download_binary(
                session, publisher_url, referer=f"https://doi.org/{doi}"
            )
            validate_download(filename, data, content_type)
            downloaded_files[filename] = {
                "data": data,
                "source_url": final_url,
                "content_type": content_type,
                "download_method": "Springer Nature publisher fallback",
            }
        except Exception as error:
            errors.append(f"{filename}: {type(error).__name__}: {error}")
    return downloaded_files, errors


def _matching_name(filename: str, expected_names: set[str]) -> bool:
    normalized_expected = {_normalized_file_name(name) for name in expected_names}
    return _normalized_file_name(filename) in normalized_expected


def _save_downloaded_file(root: Path, pmcid: str, filename: str, data: bytes) -> Path:
    safe_name = _safe_file_name(filename)
    if not safe_name:
        raise ValueError("Supplementary filename is missing")
    destination = root / pmcid / safe_name
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_bytes(data)
    return destination


def _checkpoint_path(download_root: Path, pmcid: str) -> Path:
    return download_root / pmcid / "extraction_checkpoint.json"


def _load_checkpoint(download_root: Path, pmcid: str) -> dict[str, Any] | None:
    path = _checkpoint_path(download_root, pmcid)
    if not path.is_file():
        return None
    try:
        checkpoint = json.loads(path.read_text(encoding="utf-8"))
        if (
            checkpoint.get("Checkpoint_Version") != CHECKPOINT_VERSION
            or checkpoint.get("PMCID") != pmcid
        ):
            return None
        result = checkpoint.get("Result")
        if not isinstance(result, dict):
            return None
        for file_result in result.get("Supplementary_Files", []):
            saved_path = file_result.get("Saved_Path", "")
            if saved_path and not Path(saved_path).is_file():
                return None
        return result
    except (OSError, TypeError, ValueError, json.JSONDecodeError):
        return None


def _save_checkpoint(download_root: Path, pmcid: str, result: dict[str, Any]) -> None:
    path = _checkpoint_path(download_root, pmcid)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(
        {
            "Checkpoint_Version": CHECKPOINT_VERSION,
            "PMCID": pmcid,
            "Result": result,
        },
        ensure_ascii=False,
        indent=2,
    )
    temporary_path = path.with_suffix(".tmp")
    temporary_path.write_text(payload, encoding="utf-8")
    temporary_path.replace(path)


def _project_result(
    raw_result: dict[str, Any], root: Path, pmcid: str
) -> dict[str, Any]:
    file_results: list[dict[str, Any]] = []
    for file_result in raw_result["Files"]:
        file_results.append(
            {
                "File_Name": file_result["filename"],
                "Source": file_result.get("download_method") or "",
                "Source_URL": file_result.get("source_url") or "",
                "Saved_Path": file_result.get("saved_path") or "",
                "Status": file_result.get("extraction_status") or "failed",
                "Content_Type": file_result.get("content_type") or "",
                "Byte_Size": file_result.get("byte_size", 0),
                "Text_Characters": file_result.get("text_chars", 0),
                "Error": file_result.get("error"),
            }
        )

    status_key = raw_result["Status"]
    total = len(file_results)
    extracted = sum(item["Status"] == "success" for item in file_results)
    if status_key == "success":
        status = f"Success: {extracted}/{total} files"
    elif status_key == "files_found_no_extractable_text":
        status = f"Downloaded: {total} files; no extractable text"
    elif status_key == "supplement_detected_download_failed":
        status = "Failed: supplementary material detected but download failed"
    else:
        status = "No supplementary files detected"

    return {
        "Supplementary_Content": _truncate_content(raw_result["Supplementary_Content"]),
        "Supplementary_Status": status,
        "Supplementary_Files": file_results,
    }


def _extract_supplementary_material(
    pmcid: str,
    session: requests.Session,
    root: Path,
    xml_content: bytes | str | None = None,
) -> dict[str, Any]:
    errors: list[str] = []
    html_links: list[dict[str, str]] = []
    downloaded_files: dict[str, dict[str, Any]] = {}
    html = ""
    article_url = f"https://pmc.ncbi.nlm.nih.gov/articles/{pmcid}/"

    xml_files: list[dict[str, str]] = []
    publisher_doi: str | None = None
    try:
        if xml_content is None:
            xml_content, _ = download_pmc_xml(session, pmcid)
        xml_detected = detect_supplementary_in_xml(xml_content)
        xml_files = find_supplementary_files(xml_content)
        publisher_doi = extract_doi_from_pmc_xml(xml_content)
    except Exception as error:
        xml_detected = False
        errors.append(f"XML: {type(error).__name__}: {error}")

    try:
        html, article_url = download_pmc_html(session, pmcid)
        html_links = extract_supplementary_links_from_html(html, article_url)
    except Exception as error:
        errors.append(f"PMC HTML: {type(error).__name__}: {error}")

    # PMC sometimes returns a browser-check page with no article links. The
    # normal AI_extract workflow already supplies authoritative NCBI XML, so
    # retain the notebook's URL shape using only filenames declared there.
    if not html_links and xml_files:
        pmcid_number = pmcid.removeprefix("PMC")
        html_links = [
            {
                "url": (
                    "https://pmc.ncbi.nlm.nih.gov/articles/instance/"
                    f"{pmcid_number}/bin/{quote(item['File_Name'], safe='._-')}"
                ),
                "filename": item["File_Name"],
                "link_text": "",
            }
            for item in xml_files
        ]

    expected_filenames = {item["filename"] for item in html_links}
    for link in html_links:
        filename = link["filename"]
        try:
            data, final_url, content_type = download_binary(
                session, link["url"], referer=article_url
            )
            validate_download(filename, data, content_type)
            downloaded_files[filename] = {
                "data": data,
                "source_url": final_url,
                "content_type": content_type,
                "download_method": "PMC HTML link",
            }
        except Exception as error:
            errors.append(f"PMC download {filename}: {type(error).__name__}: {error}")

    missing_files = {
        name
        for name in expected_filenames
        if not _matching_name(name, set(downloaded_files))
    }
    if not downloaded_files or missing_files:
        try:
            europe_files, europe_error = download_europe_pmc_zip(session, pmcid)
            if europe_error:
                errors.append(f"Europe PMC: {europe_error}")
            for filename, file_data in europe_files.items():
                if _matching_name(filename, set(downloaded_files)):
                    continue
                if expected_filenames and not _matching_name(filename, expected_filenames):
                    continue
                try:
                    validate_download(filename, file_data["data"], file_data["content_type"])
                    file_data["download_method"] = "Europe PMC ZIP"
                    downloaded_files[filename] = file_data
                except Exception as error:
                    errors.append(
                        f"Europe PMC file {filename}: {type(error).__name__}: {error}"
                    )
        except Exception as error:
            errors.append(f"Europe PMC: {type(error).__name__}: {error}")

    missing_files = {
        name
        for name in expected_filenames
        if not _matching_name(name, set(downloaded_files))
    }
    if not downloaded_files or missing_files:
        try:
            biostudies_files, biostudies_errors = download_biostudies_files(session, pmcid)
            errors.extend(f"BioStudies: {error}" for error in biostudies_errors)
            for filename, file_data in biostudies_files.items():
                if _matching_name(filename, set(downloaded_files)):
                    continue
                if expected_filenames and not _matching_name(filename, expected_filenames):
                    continue
                file_data["download_method"] = "BioStudies fallback"
                downloaded_files[filename] = file_data
        except Exception as error:
            errors.append(f"BioStudies: {type(error).__name__}: {error}")

    # The notebook's Springer patch is only used when the normal workflow did
    # not extract content and PMC exposed concrete supplementary links.
    if html_links and not downloaded_files:
        publisher_files, publisher_errors = download_springer_supplementary_files(
            session,
            publisher_doi or extract_doi_from_pmc_html(html),
            [item["url"] for item in html_links],
        )
        errors.extend(f"Publisher: {error}" for error in publisher_errors)
        downloaded_files.update(publisher_files)

    file_results: list[dict[str, Any]] = []
    supplementary_sections: list[str] = []
    for filename, file_data in downloaded_files.items():
        saved_path = ""
        try:
            saved_path = str(
                _save_downloaded_file(root, pmcid, filename, file_data["data"])
            )
            extracted_text = extract_text_from_file(filename, file_data["data"])
            extracted_text = _truncate_content(extracted_text)
            extraction_status = (
                "success" if extracted_text else "downloaded_no_extractable_text"
            )
            if extracted_text:
                supplementary_sections.append(
                    f"===== {filename} =====\n\n{extracted_text}"
                )
            error_message = None
        except Exception as error:
            extracted_text = ""
            extraction_status = "failed"
            error_message = f"{type(error).__name__}: {error}"

        file_results.append(
            {
                "filename": filename,
                "source_url": file_data.get("source_url", ""),
                "download_method": file_data.get("download_method"),
                "content_type": file_data.get("content_type", ""),
                "byte_size": len(file_data["data"]),
                "extraction_status": extraction_status,
                "text_chars": len(extracted_text),
                "error": error_message,
                "saved_path": saved_path,
            }
        )

    supplementary_content = "\n\n".join(supplementary_sections).strip()
    supplementary_exists = bool(xml_detected or html_links or downloaded_files)
    if supplementary_content:
        status = "success"
    elif downloaded_files:
        status = "files_found_no_extractable_text"
    elif supplementary_exists:
        status = "supplement_detected_download_failed"
    else:
        status = "no_supplementary_material"

    return {
        "PMCID": pmcid,
        "Status": status,
        "Supplementary_Exists": supplementary_exists,
        "XML_Supplement_Detected": xml_detected,
        "HTML_Links_Found": len(html_links),
        "Supplementary_Links": [item["url"] for item in html_links],
        "Files_Found": len(file_results),
        "Files": file_results,
        "Supplementary_Content": supplementary_content,
        "Errors": errors,
    }


def extract_pmc_supplementary_material(
    pmcid: Any,
    xml_content: bytes | str | None = None,
    download_root: str | Path | None = None,
    force_refresh: bool = False,
) -> dict[str, Any]:
    """Run the notebook workflow without changing the AI_extract API contract."""
    normalized_pmcid = _normalize_pmcid(pmcid)
    if not normalized_pmcid:
        return {
            "Supplementary_Content": "",
            "Supplementary_Status": "PMCID missing or invalid",
            "Supplementary_Files": [],
        }

    root = Path(download_root) if download_root else DEFAULT_DOWNLOAD_ROOT
    if not force_refresh:
        checkpoint_result = _load_checkpoint(root, normalized_pmcid)
        if checkpoint_result is not None:
            return checkpoint_result

    raw_result = _extract_supplementary_material(
        normalized_pmcid,
        _create_session(),
        root,
        xml_content=xml_content,
    )
    result = _project_result(raw_result, root, normalized_pmcid)
    _save_checkpoint(root, normalized_pmcid, result)
    return result
